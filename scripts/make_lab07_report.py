"""Build Lab 07 report (23CSE301) as Word + PDF from lab07_metrics.json.

Reads scripts/lab07_metrics.json + figures, writes:
  docs/Lab07_23CSE301_Report.docx  (python-docx)
  docs/Lab07_23CSE301_Report.pdf   (reportlab)
Run: python scripts/make_lab07_report.py
"""

from __future__ import annotations

import json
from pathlib import Path

ROOT = Path(__file__).resolve().parent.parent
SCRIPTS = ROOT / "scripts"
DOCS = ROOT / "docs"


def load_inputs():
    metrics = json.loads((SCRIPTS / "lab07_metrics.json").read_text())
    tree_fig = SCRIPTS / metrics["tree_fig"]
    boundary_fig = SCRIPTS / metrics["boundary_fig"]
    assert tree_fig.exists() and boundary_fig.exists(), "run lab07_decision_tree.py first"
    return metrics, tree_fig, boundary_fig


def build_docx(metrics, tree_fig, boundary_fig, out_path: Path):
    from docx import Document
    from docx.shared import Pt

    doc = Document()
    style = doc.styles["Normal"]
    style.font.size = Pt(11)

    doc.add_heading("23CSE301 — Lab Session 07: Decision Trees on Project Data", 0)
    doc.add_paragraph(
        "Project: Audelle — AI-powered audio editor. Dataset: GTZAN subset "
        f"({metrics['clips']} clips, genres: {', '.join(metrics['genres'])}), "
        "same Hugging Face mirror as scripts/train_genre_classifier.py. "
        "Features per clip (librosa): spectral centroid mean, zero-crossing-rate "
        "mean, RMS mean, spectral-flatness mean, tempo. Target: genre label."
    )

    doc.add_heading("1. Introduction (project understanding)", 1)
    doc.add_paragraph(
        "Audelle edits audio from natural-language prompts. Before any edit, "
        "supervised classifiers read the file: a RandomForest groove classifier "
        "(15 EDM feels), a RandomForest genre classifier (10 GTZAN genres), and a "
        "RandomForest noise-condition classifier (clean/hiss/hum/crowd/clipped/"
        "muffled) plus an SNR regressor. This lab studies the foundation of all "
        "three: the decision tree — entropy, Gini impurity, information gain, "
        "binning of continuous audio features, tree construction, and "
        "hyper-parameter tuning — applied to our own genre data."
    )

    doc.add_heading("2. Literature survey", 1)
    doc.add_paragraph(
        "scikit-learn decision-tree documentation (plot_tree) defines the "
        "splitting criteria used here: entropy and Gini impurity, with "
        "information gain selecting each split. The DataCamp decision-tree "
        "tutorial covers the classification pipeline (train/test split, "
        "visualization, decision boundaries) followed in A6–A7. The sklearn "
        "iris decision-tree example is the template for the 2-feature boundary "
        "plot. In music-information retrieval, GTZAN genre classification from "
        "spectral features (MFCC, centroid, ZCR) is the standard benchmark our "
        "project's genre head follows."
    )

    doc.add_heading("3. Methodology (A1–A8)", 1)
    doc.add_paragraph(
        "A1: dataset entropy H = -sum p_i log2 p_i over the 4 genre labels "
        "(equal_width_binning provided for continuous outcomes). A2: Gini = "
        "1 - sum p_j^2. A3: each continuous feature binned (3 bins), "
        "information gain = parent entropy − weighted child entropy; highest "
        "gain becomes the root. A4: bin_feature(values, n_bins=4, "
        "binning_type='equal_width'|'equal_frequency') with defaults. A5: "
        "hand-rolled recursive tree (multi-way splits on binned features, "
        "majority fallback, max_depth=3). A6: sklearn DecisionTreeClassifier "
        "visualized with plot_tree. A7: 2-feature (centroid, ZCR) tree with "
        "meshgrid contour boundary. A8: GridSearchCV over criterion, "
        "max_depth, min_samples_split (cv=3)."
    )

    doc.add_heading("4. Results", 1)
    doc.add_paragraph(
        f"Dataset entropy: {metrics['entropy']} bits (maximum for 4 balanced "
        f"classes). Gini: {metrics['gini']}. Root node by information gain: "
        f"{metrics['root_feature']}."
    )
    table = doc.add_table(rows=1, cols=2)
    table.style = "Light Grid Accent 1"
    table.rows[0].cells[0].text = "Feature"
    table.rows[0].cells[1].text = "Information gain"
    for name, gain in sorted(metrics["gains"].items(), key=lambda kv: -kv[1]):
        row = table.add_row().cells
        row[0].text = name
        row[1].text = str(gain)
    doc.add_paragraph(
        f"Hand-rolled tree test accuracy: {metrics['own_tree_test_accuracy']}. "
        f"Tuned classifier: {metrics['best_params']} "
        f"(CV score {metrics['best_cv_score']})."
    )
    doc.add_paragraph("Decision tree (sklearn, max_depth=3):")
    doc.add_picture(str(tree_fig), width=Pt(430))
    doc.add_paragraph("Decision boundary on centroid vs ZCR:")
    doc.add_picture(str(boundary_fig), width=Pt(380))

    doc.add_heading("5. Observations, inferences, discussion", 1)
    doc.add_paragraph(
        "Spectral centroid dominates the root split (gain "
        f"{metrics['gains']['centroid_mean']}), which matches acoustics: metal "
        "is bright, classical is dark, so one centroid threshold already "
        "separates half the classes; tempo carries almost no signal "
        f"(gain {metrics['gains']['tempo']}) because all four genres overlap "
        "in BPM. The hand-rolled tree scores lower than the tuned sklearn "
        "tree: multi-way splits on 3-bin features fragment 24 samples into "
        "tiny leaves, while sklearn's binary splits with min_samples_split "
        "regularize better — expected behavior on small data, and the reason "
        "the product uses forests (bagging over many such trees). The "
        "boundary plot shows axis-aligned rectangles, the visual signature of "
        "tree classifiers. GridSearchCV selects a shallow depth-3 gini tree, "
        "confirming that capacity control matters more than criterion choice "
        "here."
    )

    doc.add_heading("6. Conclusion", 1)
    doc.add_paragraph(
        "Entropy/Gini quantified a balanced 4-class problem; information gain "
        "identified spectral centroid as the root; a hand-built tree, a "
        "visualized sklearn tree, a 2-D boundary, and GridSearchCV tuning "
        "complete the pipeline on our own GTZAN project data. The experiment "
        "directly justifies the product's RandomForest heads, which are "
        "ensembles of exactly these tuned trees."
    )
    doc.save(out_path)
    return str(out_path)


def build_pdf(metrics, tree_fig, boundary_fig, out_path: Path):
    from reportlab.lib.pagesizes import A4
    from reportlab.lib.styles import getSampleStyleSheet
    from reportlab.lib.units import cm
    from reportlab.platypus import Image, Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle

    styles = getSampleStyleSheet()
    story = []
    add = story.append
    add(Paragraph("23CSE301 — Lab Session 07: Decision Trees on Project Data", styles["Title"]))
    add(Spacer(1, 0.3 * cm))
    add(Paragraph(
        "Project: Audelle — AI-powered audio editor. Dataset: GTZAN subset "
        f"({metrics['clips']} clips, genres: {', '.join(metrics['genres'])}), "
        "same Hugging Face mirror as scripts/train_genre_classifier.py. "
        "Features per clip (librosa): spectral centroid mean, ZCR mean, RMS "
        "mean, spectral-flatness mean, tempo. Target: genre label.",
        styles["Normal"]))
    add(Paragraph("1. Introduction (project understanding)", styles["Heading1"]))
    add(Paragraph(
        "Audelle edits audio from natural-language prompts. Before any edit, "
        "supervised classifiers read the file: a RandomForest groove classifier "
        "(15 EDM feels), a RandomForest genre classifier (10 GTZAN genres), and a "
        "RandomForest noise-condition classifier plus SNR regressor. This lab "
        "studies their foundation — the decision tree — on our own genre data.",
        styles["Normal"]))
    add(Paragraph("2. Literature survey", styles["Heading1"]))
    add(Paragraph(
        "scikit-learn plot_tree docs (entropy/Gini criteria), the DataCamp "
        "decision-tree tutorial (classification pipeline, boundaries), the "
        "sklearn iris example (2-D boundary template), and GTZAN spectral "
        "genre classification as the MIR benchmark behind our genre head.",
        styles["Normal"]))
    add(Paragraph("3. Methodology (A1–A8)", styles["Heading1"]))
    add(Paragraph(
        "A1 entropy H = -sum p log2 p; A2 Gini = 1 - sum p^2; A3 root by "
        "information gain on binned features; A4 bin_feature with "
        "binning_type/n_bins defaults; A5 hand-rolled recursive multi-way "
        "tree (max_depth=3); A6 plot_tree figure; A7 centroid-vs-ZCR meshgrid "
        "boundary; A8 GridSearchCV over criterion/depth/split, cv=3.",
        styles["Normal"]))
    add(Paragraph("4. Results", styles["Heading1"]))
    add(Paragraph(
        f"Entropy: {metrics['entropy']} bits; Gini: {metrics['gini']}; "
        f"root: {metrics['root_feature']}; hand-rolled test accuracy: "
        f"{metrics['own_tree_test_accuracy']}; tuned: {metrics['best_params']} "
        f"(CV {metrics['best_cv_score']}).", styles["Normal"]))
    rows = [["Feature", "Information gain"]] + [
        [k, str(v)] for k, v in sorted(metrics["gains"].items(), key=lambda kv: -kv[1])]
    table = Table(rows)
    table.setStyle(TableStyle([("GRID", (0, 0), (-1, -1), 0.5, (0, 0, 0))]))
    add(table)
    add(Spacer(1, 0.3 * cm))
    add(Paragraph("Decision tree (sklearn, max_depth=3):", styles["Normal"]))
    add(Image(str(tree_fig), width=15 * cm, height=8.6 * cm))
    add(Spacer(1, 0.3 * cm))
    add(Paragraph("Decision boundary (centroid vs ZCR):", styles["Normal"]))
    add(Image(str(boundary_fig), width=13 * cm, height=9.3 * cm))
    add(Paragraph("5. Observations, inferences, discussion", styles["Heading1"]))
    add(Paragraph(
        "Centroid dominates the root split "
        f"(gain {metrics['gains']['centroid_mean']}); tempo is near-useless "
        f"(gain {metrics['gains']['tempo']}) as genres overlap in BPM. The "
        "hand-rolled tree trails sklearn: multi-way binned splits fragment 24 "
        "samples while binary splits with min_samples_split regularize — why "
        "the product uses forests. Rectangular boundary regions are the "
        "tree signature. Shallow depth-3 wins tuning: capacity control beats "
        "criterion choice here.", styles["Normal"]))
    add(Paragraph("6. Conclusion", styles["Heading1"]))
    add(Paragraph(
        "Balanced 4-class problem quantified, centroid root identified, "
        "hand-built + visualized + boundary + tuned pipeline completed on "
        "project GTZAN data — directly justifying the product's RandomForest "
        "heads as ensembles of these tuned trees.", styles["Normal"]))
    SimpleDocTemplate(str(out_path), pagesize=A4).build(story)
    return str(out_path)


def main():
    metrics, tree_fig, boundary_fig = load_inputs()
    docx_path = build_docx(metrics, tree_fig, boundary_fig,
                           DOCS / "Lab07_23CSE301_Report.docx")
    print(f"Word: {docx_path}")
    pdf_path = build_pdf(metrics, tree_fig, boundary_fig,
                         DOCS / "Lab07_23CSE301_Report.pdf")
    print(f"PDF: {pdf_path}")


if __name__ == "__main__":
    main()
